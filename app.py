import os
import io
import sqlite3
import secrets
from datetime import datetime
from functools import wraps

import requests
from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request, redirect, url_for, session, send_file, abort
from werkzeug.security import check_password_hash

load_dotenv()
app = Flask(__name__)
secret_key = os.getenv("FLASK_SECRET_KEY")
if not secret_key:
    raise RuntimeError("FLASK_SECRET_KEY is not set. Add a long random secret in environment variables.")
app.secret_key = secret_key
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax", SESSION_COOKIE_SECURE=os.getenv("COOKIE_SECURE", "false").lower() == "true")
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
DB_PATH = os.getenv("DATABASE_PATH", os.path.join(os.path.dirname(__file__), "bookings.sqlite3"))
IS_POSTGRES = DATABASE_URL.startswith(("postgres://", "postgresql://"))
ADMIN_USER = os.getenv("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD_HASH = os.getenv("ADMIN_PASSWORD_HASH", "")
TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN", "")
TELEGRAM_CHAT_ID = os.getenv("TELEGRAM_CHAT_ID", "")
ALLOWED_TIMES = {"10:00", "11:00", "12:00", "13:00", "14:00", "15:00", "16:00"}

if IS_POSTGRES:
    try:
        import psycopg
        from psycopg.rows import dict_row
    except ImportError as exc:
        raise RuntimeError("Install psycopg[binary] to use PostgreSQL") from exc

def connect_db():
    if IS_POSTGRES:
        return psycopg.connect(DATABASE_URL, row_factory=dict_row, connect_timeout=10)
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con

def placeholder():
    return "%s" if IS_POSTGRES else "?"

def init_db():
    with connect_db() as con:
        if IS_POSTGRES:
            con.execute("""CREATE TABLE IF NOT EXISTS bookings (
                id BIGSERIAL PRIMARY KEY,
                created_at VARCHAR(32) NOT NULL,
                name VARCHAR(100) NOT NULL,
                phone VARCHAR(40) NOT NULL,
                date VARCHAR(10) NOT NULL,
                time VARCHAR(5) NOT NULL,
                status VARCHAR(30) NOT NULL DEFAULT 'Новая',
                note TEXT NOT NULL DEFAULT ''
            )""")
        else:
            con.execute("""CREATE TABLE IF NOT EXISTS bookings (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                name TEXT NOT NULL,
                phone TEXT NOT NULL,
                date TEXT NOT NULL,
                time TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'Новая',
                note TEXT NOT NULL DEFAULT ''
            )""")

init_db()

def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("admin_logged_in"):
            return redirect(url_for("login", next=request.path))
        return fn(*args, **kwargs)
    return wrapper

def send_telegram(booking):
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        return False, "Telegram не настроен"
    message = ("🪄 Новая заявка на квест «Гарри Поттер» в Экспериментарии\n\n"
        f"№: {booking['id']}\nИмя: {booking['name']}\nТелефон: {booking['phone']}\n"
        f"Дата: {booking['date']}\nВремя: {booking['time']}\nСоздана: {booking['created_at']}")
    try:
        response = requests.post(f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendMessage",
            json={"chat_id": TELEGRAM_CHAT_ID, "text": message}, timeout=8)
        response.raise_for_status()
        if not response.json().get("ok"):
            return False, "Telegram API не подтвердил отправку"
        return True, "ok"
    except Exception as exc:
        app.logger.warning("Telegram notification failed: %s", exc)
        return False, "Не удалось отправить уведомление в Telegram"

@app.get("/")
def home():
    return render_template("index.html")

@app.post("/api/bookings")
def create_booking():
    data = request.get_json(silent=True) or {}
    name = str(data.get("name", "")).strip()
    phone = str(data.get("phone", "")).strip()
    date_value = str(data.get("date", "")).strip()
    time_value = str(data.get("time", "")).strip()
    if not name or len(name) > 100:
        return jsonify(error="Укажите имя (до 100 символов)."), 400
    digits = "".join(ch for ch in phone if ch.isdigit())
    if len(digits) < 10 or len(digits) > 15 or len(phone) > 40:
        return jsonify(error="Укажите корректный номер телефона."), 400
    try:
        date_obj = datetime.strptime(date_value, "%Y-%m-%d").date()
        if date_obj < datetime.now().date():
            return jsonify(error="Выберите сегодняшнюю или будущую дату."), 400
    except ValueError:
        return jsonify(error="Выберите корректную дату."), 400
    if time_value not in ALLOWED_TIMES:
        return jsonify(error="Выберите время с 10:00 до 16:00."), 400
    created_at = datetime.now().strftime("%d.%m.%Y %H:%M")
    ph = placeholder()
    with connect_db() as con:
        if IS_POSTGRES:
            row = con.execute(f"INSERT INTO bookings(created_at,name,phone,date,time) VALUES({ph},{ph},{ph},{ph},{ph}) RETURNING *",
                (created_at, name, phone, date_obj.strftime("%d.%m.%Y"), time_value)).fetchone()
        else:
            cur = con.execute(f"INSERT INTO bookings(created_at,name,phone,date,time) VALUES({ph},{ph},{ph},{ph},{ph})",
                (created_at, name, phone, date_obj.strftime("%d.%m.%Y"), time_value))
            row = con.execute("SELECT * FROM bookings WHERE id=?", (cur.lastrowid,)).fetchone()
        booking = dict(row)
    send_telegram(booking)
    return jsonify(ok=True, id=booking["id"]), 201

@app.get("/login")
def login():
    if session.get("admin_logged_in"):
        return redirect(url_for("admin"))
    return render_template("login.html", error=request.args.get("error", ""), next=request.args.get("next", ""))

@app.post("/login")
def login_post():
    username = request.form.get("username", "")
    password = request.form.get("password", "")
    valid_password = bool(ADMIN_PASSWORD_HASH) and check_password_hash(ADMIN_PASSWORD_HASH, password)
    if secrets.compare_digest(username, ADMIN_USER) and valid_password:
        session.clear()
        session["admin_logged_in"] = True
        return redirect(request.form.get("next") or url_for("admin"))
    return redirect(url_for("login", error="Неверный логин или пароль"))

@app.post("/logout")
@admin_required
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.get("/admin")
@admin_required
def admin():
    status_filter = request.args.get("status", "")
    ph = placeholder()
    with connect_db() as con:
        if status_filter:
            rows = con.execute(f"SELECT * FROM bookings WHERE status={ph} ORDER BY id DESC", (status_filter,)).fetchall()
        else:
            rows = con.execute("SELECT * FROM bookings ORDER BY id DESC").fetchall()
        total = con.execute("SELECT COUNT(*) AS n FROM bookings").fetchone()["n"]
        fresh = con.execute(f"SELECT COUNT(*) AS n FROM bookings WHERE status={ph}", ("Новая",)).fetchone()["n"]
    return render_template("admin.html", bookings=[dict(r) for r in rows], total=total, fresh=fresh, status_filter=status_filter)

@app.post("/admin/<int:booking_id>/update")
@admin_required
def update_booking(booking_id):
    status = request.form.get("status", "Новая")
    note = request.form.get("note", "").strip()[:1000]
    if status not in {"Новая", "Связались", "Подтверждена", "Отменена", "Завершена"}:
        abort(400)
    ph = placeholder()
    with connect_db() as con:
        con.execute(f"UPDATE bookings SET status={ph}, note={ph} WHERE id={ph}", (status, note, booking_id))
    return redirect(url_for("admin", status=request.args.get("status", "")))

@app.get("/admin/export.xlsx")
@admin_required
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    with connect_db() as con:
        rows = con.execute("SELECT id,created_at,name,phone,date,time,status,note FROM bookings ORDER BY id DESC").fetchall()
    wb = Workbook(); ws = wb.active; ws.title = "Заявки на квест"
    ws.append(["№", "Создана", "Имя", "Телефон", "Дата квеста", "Время", "Статус", "Комментарий"])
    for row in rows:
        values = [row[k] for k in ("id","created_at","name","phone","date","time","status","note")]
        ws.append(values)
        ws.cell(ws.max_row, 4).number_format = "@"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="541B2B")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for idx, width in enumerate([8,19,24,22,18,12,18,45], start=1): ws.column_dimensions[get_column_letter(idx)].width = width
    output = io.BytesIO(); wb.save(output); output.seek(0)
    filename = "zayavki_kvest_" + datetime.now().strftime("%Y-%m-%d") + ".xlsx"
    return send_file(output, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.get("/health")
def health():
    with connect_db() as con:
        con.execute("SELECT 1")
    return jsonify(ok=True, database="postgresql" if IS_POSTGRES else "sqlite")

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5000")), debug=False)
